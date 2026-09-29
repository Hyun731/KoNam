"""업로드 문서에서 텍스트를 뽑는다. PDF·DOCX·TXT는 로컬에서, 이미지·스캔 PDF는 OpenAI로 읽는다.

쪽 경계는 폼피드(\\f)로 남긴다. 조항 분할(segment.py)이 이를 읽어 조항마다 쪽 번호를 붙인다.
"""

import io
from dataclasses import dataclass

from app import llm

MAX_BYTES = 15 * 1024 * 1024
MAX_CHARS = 80_000
PAGE_BREAK = "\f"

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
IMAGE_TYPES = {"image/png", "image/jpeg", "image/webp", "image/heic"}

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


class UnsupportedDocument(ValueError):
    pass


@dataclass
class Extracted:
    text: str  # 쪽 사이는 \f
    page_count: int | None
    method: str  # pdf-text | docx | txt | ocr-llm


def count_pages(text: str) -> int:
    return text.count(PAGE_BREAK) + 1


def _mime(filename: str, mime: str | None) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        return PDF
    if name.endswith(".docx"):
        return DOCX
    if name.endswith(".txt"):
        return "text/plain"
    if name.endswith((".png", ".jpg", ".jpeg", ".webp", ".heic")):
        return mime if mime in IMAGE_TYPES else "image/jpeg"
    if name.endswith((".hwp", ".hwpx", ".doc")):
        raise UnsupportedDocument("HWP and DOC files are not supported yet. Please save it as PDF or DOCX, "
                                  "or paste the text instead.")
    return mime or "application/octet-stream"


def _docx_text(data: bytes) -> str:
    """본문을 문서 순서대로 읽는다. 명시적 쪽 나누기(w:br type=page)는 \\f로 바꾼다."""
    from docx import Document

    doc = Document(io.BytesIO(data))
    out: list[str] = []
    for el in doc.element.body.iterchildren():
        if el.tag == _W + "p":
            line = []
            for node in el.iter(_W + "t", _W + "br", _W + "tab"):
                if node.tag == _W + "t":
                    line.append(node.text or "")
                elif node.tag == _W + "tab":
                    line.append("\t")
                elif node.get(_W + "type") == "page":
                    out.append("".join(line))
                    out.append(PAGE_BREAK)
                    line = []
                else:
                    line.append("\n")
            out.append("".join(line))
        elif el.tag == _W + "tbl":
            for row in el.iter(_W + "tr"):
                cells = ["".join(t.text or "" for t in c.iter(_W + "t")).strip() for c in row.iter(_W + "tc")]
                out.append("\t".join(cells))
    text = "\n".join(out)
    return text.replace("\n" + PAGE_BREAK + "\n", PAGE_BREAK)


def _clip(text: str) -> str:
    return text.strip()[:MAX_CHARS]


async def extract_text(filename: str, mime: str | None, data: bytes) -> Extracted:
    if len(data) > MAX_BYTES:
        raise UnsupportedDocument("The file is too large (max 15 MB).")
    kind = _mime(filename, mime)

    if kind == PDF:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        pages = len(reader.pages)
        text = PAGE_BREAK.join((p.extract_text() or "").strip() for p in reader.pages)
        if len(text.replace(PAGE_BREAK, "").strip()) < 80:  # 스캔 PDF → 모델로 읽기 (쪽 구분은 알 수 없음)
            text = await llm.read_document(data, PDF, filename)
            return Extracted(_clip(text), pages, "ocr-llm")
        return Extracted(_clip(text), pages, "pdf-text")

    if kind == DOCX:
        text = _clip(_docx_text(data))
        return Extracted(text, count_pages(text), "docx")

    if kind == "text/plain":
        for enc in ("utf-8", "utf-16", "cp1258", "latin-1"):
            try:
                text = _clip(data.decode(enc))
            except UnicodeDecodeError:
                continue
            return Extracted(text, count_pages(text), "txt")

    if kind in IMAGE_TYPES:
        text = await llm.read_document(data, kind, filename)
        return Extracted(_clip(text).replace(PAGE_BREAK, "\n"), 1, "ocr-llm")

    raise UnsupportedDocument("Only PDF, DOCX, TXT and image (JPG/PNG) files are supported.")
