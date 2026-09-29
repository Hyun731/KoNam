import asyncio
import io

from docx import Document

from app.documents.extract import PAGE_BREAK, count_pages, extract_text
from app.documents.segment import segment_document


def test_segments_carry_page_numbers_across_form_feeds():
    text = (
        "HỢP ĐỒNG THUÊ XE\nBên A: Công ty X\n"
        "Điều 1. Đối tượng\nXe mô tô biển số 59X1.\n"
        + PAGE_BREAK
        + "Điều 2. Giá thuê\nGiá thuê 200.000 đồng/ngày.\nTiếp tục sang trang sau\n"
        + PAGE_BREAK
        + "vẫn thuộc Điều 2.\nĐiều 3. Đặt cọc\nĐặt cọc 5.000.000 đồng."
    )
    segs = segment_document(text)
    by_title = {s.title.split(".")[0]: s for s in segs}
    assert segs[0].pages == [1]  # 머리말
    assert by_title["Điều 1"].pages == [1]
    assert by_title["Điều 2"].pages == [2, 3]
    assert by_title["Điều 3"].pages == [3]
    assert all(PAGE_BREAK not in s.text for s in segs)
    assert count_pages(text) == 3


def test_plain_text_without_form_feed_is_page_one():
    segs = segment_document("Một đoạn văn ngắn về vi phạm giao thông.\nDòng thứ hai.")
    assert [s.pages for s in segs] == [[1]]


def test_blank_leading_page_is_not_counted():
    segs = segment_document("\n" + PAGE_BREAK + "Nội dung ở trang hai của văn bản này.")
    assert segs[0].pages == [2]


def test_docx_explicit_page_break_becomes_new_page():
    doc = Document()
    doc.add_paragraph("Điều 1. Trang một")
    doc.add_page_break()
    doc.add_paragraph("Điều 2. Trang hai")
    t = doc.add_table(rows=1, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Tiền phạt", "5.000.000 đồng"
    buf = io.BytesIO()
    doc.save(buf)
    ex = asyncio.run(extract_text("a.docx", None, buf.getvalue()))
    assert ex.page_count == 2
    first, second = ex.text.split(PAGE_BREAK)
    assert "Trang một" in first and "Trang hai" in second and "Tiền phạt\t5.000.000 đồng" in second


def test_txt_form_feed_page_count():
    ex = asyncio.run(extract_text("a.txt", None, "trang một\ftrang hai\ftrang ba".encode()))
    assert ex.page_count == 3 and ex.method == "txt"
